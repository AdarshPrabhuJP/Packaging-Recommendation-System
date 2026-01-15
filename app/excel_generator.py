from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
import io

def generate_excel_report(analytics_data, materials_df):
    buffer = io.BytesIO()
    wb = Workbook()
    
    ws_summary = wb.active
    ws_summary.title = "Summary"
    
    ws_summary['A1'] = "Packaging Recommendation System - Analytics Report"
    ws_summary['A1'].font = Font(size=14, bold=True, color="FFFFFF")
    ws_summary['A1'].fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    ws_summary.merge_cells('A1:B1')
    
    ws_summary['A3'] = "Generated Date:"
    ws_summary['B3'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    
    ws_summary['A5'] = "Key Metrics"
    ws_summary['A5'].font = Font(bold=True, size=12)
    
    key_metrics = analytics_data['key_metrics']
    metrics_data = [
        ['Total Materials', key_metrics['total_materials']],
        ['CO₂ Saved (%)', f"{key_metrics['co2_saved_percent']}%"],
        ['Cost Saved (₹)', f"₹{key_metrics['cost_saved']:.2f}"],
        ['Eco-Friendly (%)', f"{key_metrics['eco_friendly_percent']}%"]
    ]
    
    row = 6
    for metric in metrics_data:
        ws_summary[f'A{row}'] = metric[0]
        ws_summary[f'B{row}'] = metric[1]
        ws_summary[f'A{row}'].font = Font(bold=True)
        row += 1
    
    for col in ['A', 'B']:
        ws_summary.column_dimensions[col].width = 25
    
    ws_materials = wb.create_sheet("Materials Data")
    
    headers = ['ID', 'Material Type', 'Cost (₹)', 'Durability', 'CO₂ (kg)', 
               'Biodegradability', 'Recyclable', 'Recommended Use']
    ws_materials.append(headers)
    
    for cell in ws_materials[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
        cell.alignment = Alignment(horizontal='center')
    
    for _, row in materials_df.iterrows():
        ws_materials.append([
            row['id'],
            row['material_type'],
            float(row['cost_per_unit']),
            row['durability_score'],
            float(row['co2_footprint']),
            row['biodegradability_score'],
            'Yes' if row['recyclable'] else 'No',
            row['recommended_use']
        ])
    
    for col in range(1, len(headers) + 1):
        ws_materials.column_dimensions[get_column_letter(col)].width = 20
    
    ws_analytics = wb.create_sheet("Analytics")
    
    ws_analytics['A1'] = "CO₂ Reduction Analysis"
    ws_analytics['A1'].font = Font(bold=True, size=12)
    
    co2_data = analytics_data['co2_reduction']
    ws_analytics['A2'] = "Baseline CO₂ (kg)"
    ws_analytics['B2'] = co2_data['baseline_co2']
    ws_analytics['A3'] = "Eco Materials CO₂ (kg)"
    ws_analytics['B3'] = co2_data['eco_co2']
    ws_analytics['A4'] = "Reduction (%)"
    ws_analytics['B4'] = f"{co2_data['reduction_percent']}%"
    
    ws_analytics['A6'] = "Cost Savings Analysis"
    ws_analytics['A6'].font = Font(bold=True, size=12)
    
    cost_data = analytics_data['cost_savings']
    ws_analytics['A7'] = "Baseline Cost (₹)"
    ws_analytics['B7'] = cost_data['baseline_cost']
    ws_analytics['A8'] = "Budget Cost (₹)"
    ws_analytics['B8'] = cost_data['budget_cost']
    ws_analytics['A9'] = "Savings (%)"
    ws_analytics['B9'] = f"{cost_data['savings_percent']}%"
    
    for col in ['A', 'B']:
        ws_analytics.column_dimensions[col].width = 25
    
    wb.save(buffer)
    buffer.seek(0)
    return buffer
